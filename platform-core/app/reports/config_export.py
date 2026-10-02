"""Xuất cấu hình máy ra Excel 3 sheet — y hệt nút "Xuất cấu hình" của hệ cũ (CHỈ ĐỌC: D:\\NetworkInventory\\app.py
`/export_config`): Tổng hợp theo phòng · Nhóm cấu hình (Model + CPU + RAM + Ổ cứng) · Chi tiết máy.
Chỉ gồm phòng CHÍNH THỨC (bỏ Phòng chờ / Khu vực … như hệ cũ). Hàm thuần: nhận danh sách máy, trả bytes .xlsx.
"""
from __future__ import annotations

import re
from collections import Counter
from collections.abc import Iterable
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.legacy_adapter.identity import is_official_room


def _nat(s: object) -> list:
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(s or ""))]


def _ram_total(s: object) -> str:
    m = re.match(r"\s*(\d+(?:\.\d+)?)\s*GB", str(s or ""), re.I)
    return (m.group(1) + "GB") if m else (str(s or "").strip() or "N/A")


def _disk(s: object) -> str:
    return re.sub(r"^\s*\d+\s*ổ:\s*", "", str(s or "")).strip() or "N/A"


def _summ(items: Iterable[str]) -> str:
    return "; ".join(f"{k} × {v}" for k, v in Counter(items).most_common())


def build(devices: list[dict], now: datetime | None = None) -> tuple[bytes, int]:
    """`devices` = kết quả devices.service.public (có room_name, hostname, is_online…). Trả (bytes .xlsx, số máy)."""
    now = now or datetime.now()
    machines = []
    for d in devices:
        room = str(d.get("room_name") or "")
        if not is_official_room(room):
            continue
        machines.append({
            "room": room, "host": str(d.get("hostname") or ""), "online": bool(d.get("is_online")),
            "ip": d.get("ip") or "", "mac": d.get("mac") or "", "model": str(d.get("model") or "").strip() or "N/A",
            "serial": d.get("serial") or "", "cpu": str(d.get("cpu") or "").strip() or "N/A",
            "ram_total": _ram_total(d.get("ram")), "ram": d.get("ram") or "", "disk": _disk(d.get("disk")),
            "mainboard": d.get("mainboard") or "", "network": d.get("network") or "",
            "ver": d.get("client_version") or "", "last": d.get("last_heartbeat_at") or ""})
    machines.sort(key=lambda m: (_nat(m["room"]), _nat(m["host"])))
    by_room: dict[str, list[dict]] = {}
    for m in machines:
        by_room.setdefault(m["room"], []).append(m)

    wb = Workbook()
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    tot_fill = PatternFill("solid", fgColor="DDEBF7")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, title, headers, data, widths, total=None):
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
        ws.append([f"Xuất lúc {now:%d/%m/%Y %H:%M} — {len(machines)} máy, {len(by_room)} phòng (không gồm Phòng chờ)"])
        ws["A2"].font = Font(italic=True, color="7F7F7F")
        ws.append(headers)
        for c in ws[3]:
            c.font, c.fill, c.border = hdr_font, hdr_fill, border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in data:
            ws.append(row)
            for c in ws[ws.max_row]:
                c.border, c.alignment = border, wrap
        if total:
            ws.append(total)
            for c in ws[ws.max_row]:
                c.font, c.fill, c.border = Font(bold=True), tot_fill, border
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{ws.max_row - (1 if total else 0)}"

    ws = wb.active
    assert ws is not None
    ws.title = "Tổng hợp theo phòng"
    data = [[room, len(ms), sum(1 for m in ms if m["online"]), _summ(m["model"] for m in ms),
             _summ(m["cpu"] for m in ms), _summ(m["ram_total"] for m in ms), _summ(m["disk"] for m in ms),
             _summ(m["ver"] or "N/A" for m in ms)] for room, ms in by_room.items()]
    sheet(ws, "TỔNG HỢP CẤU HÌNH THEO PHÒNG MÁY",
          ["Phòng", "Số máy", "Đang online", "Model", "CPU", "RAM", "Ổ cứng", "Phiên bản agent"],
          data, [18, 9, 11, 34, 42, 18, 48, 22],
          total=["TOÀN TRƯỜNG", len(machines), sum(1 for m in machines if m["online"]), "", "", "", "", ""])

    ws = wb.create_sheet("Nhóm cấu hình")
    groups: dict[tuple, list[dict]] = {}
    for m in machines:
        groups.setdefault((m["model"], m["cpu"], m["ram_total"], m["disk"]), []).append(m)
    data = []
    for i, (key, ms) in enumerate(sorted(groups.items(), key=lambda kv: -len(kv[1])), 1):
        rc = Counter(m["room"] for m in ms)
        data.append([i, key[0], key[1], key[2], key[3], len(ms),
                     "; ".join(f"{r} ({n})" for r, n in sorted(rc.items(), key=lambda x: _nat(x[0])))])
    sheet(ws, "THỐNG KÊ NHÓM CẤU HÌNH", ["STT", "Model", "CPU", "RAM", "Ổ cứng", "Số máy", "Phân bổ theo phòng"],
          data, [6, 26, 38, 10, 44, 9, 60], total=["", "TỔNG", "", "", "", len(machines), ""])

    ws = wb.create_sheet("Chi tiết máy")
    data = [[i, m["room"], m["host"], "Online" if m["online"] else "Offline", m["ip"], m["mac"], m["model"],
             m["serial"], m["cpu"], m["ram_total"], m["ram"], m["disk"], m["mainboard"], m["network"], m["ver"],
             m["last"]] for i, m in enumerate(machines, 1)]
    sheet(ws, "CHI TIẾT CẤU HÌNH TỪNG MÁY",
          ["STT", "Phòng", "Tên máy", "Trạng thái", "IP", "MAC", "Model", "Service Tag", "CPU", "RAM",
           "RAM chi tiết", "Ổ cứng", "Mainboard", "Card mạng", "Agent", "Cập nhật lần cuối"],
          data, [6, 16, 13, 10, 14, 18, 22, 12, 36, 8, 48, 44, 12, 40, 10, 19])
    for row in ws.iter_rows(min_row=4, min_col=4, max_col=4):
        for c in row:
            c.font = Font(bold=True, color="198754" if c.value == "Online" else "9C0006")

    out = BytesIO()
    wb.save(out)
    return out.getvalue(), len(machines)
