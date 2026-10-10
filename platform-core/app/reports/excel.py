"""Xuất bảng ra .xlsx cùng kiểu trình bày với "Xuất cấu hình" (config_export.py): dòng tiêu đề, dòng "Xuất lúc …",
hàng tiêu đề cột nền xanh, khung mảnh, ngắt dòng, cố định hàng tiêu đề + bộ lọc. Hàm thuần — trả bytes."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass
class Sheet:
    name: str          # tên sheet (≤ 31 ký tự)
    title: str         # dòng tiêu đề lớn
    headers: list[str]
    rows: list[list]
    widths: list[int]
    note: str = ""     # nối sau "Xuất lúc …"
    total: list | None = None
    wrap_cols: set[int] = field(default_factory=set)  # cột (1-based) cần giữ xuống dòng (mô tả nhiều dòng)


def build(sheets: list[Sheet], now: datetime | None = None) -> bytes:
    now = now or datetime.now()
    wb = Workbook()
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    tot_fill = PatternFill("solid", fgColor="DDEBF7")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    first = True
    for sh in sheets:
        if first:
            ws = wb.active
            assert ws is not None
            first = False
        else:
            ws = wb.create_sheet()
        ws.title = sh.name[:31]
        ws.append([sh.title])
        ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
        ws.append([f"Xuất lúc {now:%d/%m/%Y %H:%M}" + (f" — {sh.note}" if sh.note else "")])
        ws["A2"].font = Font(italic=True, color="7F7F7F")
        ws.append(sh.headers)
        for c in ws[3]:
            c.font, c.fill, c.border = hdr_font, hdr_fill, border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sh.rows:
            ws.append(row)
            for i, c in enumerate(ws[ws.max_row], 1):
                c.border = border
                c.alignment = Alignment(wrap_text=True, vertical="top") if i in sh.wrap_cols else Alignment(
                    vertical="top")
        if sh.total:
            ws.append(sh.total)
            for c in ws[ws.max_row]:
                c.font, c.fill, c.border = Font(bold=True), tot_fill, border
        for i, w in enumerate(sh.widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A4"
        last = ws.max_row - (1 if sh.total else 0)
        if last >= 3:
            ws.auto_filter.ref = f"A3:{get_column_letter(len(sh.headers))}{last}"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def response(data: bytes, filename: str):
    """Trả tệp tải về. `filename` không dấu (header Content-Disposition chỉ nhận latin-1)."""
    from flask import make_response

    resp = make_response(data)
    resp.headers["Content-Type"] = XLSX
    resp.headers["Content-Disposition"] = f"attachment; filename={filename}"
    return resp
