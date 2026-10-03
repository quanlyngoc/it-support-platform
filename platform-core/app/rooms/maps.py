"""Sơ đồ phòng (Sổ tay 5.3) — mô hình có cấu trúc trong `rooms.map` + tương thích TỆP sơ đồ của hệ cũ.

Hệ cũ lưu sơ đồ là ẢNH CHỤP DOM của trình vẽ (`room_detail.html` captureState): `bounds` (viền phòng), `canvasWidth`,
`items[id] = {t, l, w, h, html: outerHTML, parentId}`, `pcClasses`, `machineGridHTML`. Hệ thi (ThiHSK/ThiPM,
`shared/web/sodo.js`) đọc THẲNG các tệp đó qua `MONITOR_MAPS_DIR`: lấy mọi `.pc-node` trong `items[].html`, chữ trong
`.pc-name` (số ghế, khớp với tên máy), toạ độ `l`/`t` và lớp hướng `dir-*`.

Core KHÔNG lưu HTML: lưu mô hình `MapModel` (toạ độ cùng hệ px với trình vẽ cũ — nền 1000×1300, máy 85×95), rồi mỗi lần
lưu thì XUẤT lại tệp đúng định dạng cũ (`to_legacy`) để hệ thi đọc như trước. Phòng chưa có sơ đồ trong Core thì đọc tệp
của hệ cũ (`from_legacy`, CHỈ ĐỌC kho cũ).

Quy tắc nhập lưới Excel giữ đúng hệ cũ (`renderGrid`, `mapColToX`, `mapRowToY`): máy 85×95, 2 máy một bàn cách 20,
giữa 2 bàn 80, hàng cách 22, gốc (90, 90); số 1 chữ số thêm 0 phía trước; dòng đầu là tiêu đề.
"""
from __future__ import annotations

import io
import json
import os
import re
import tempfile
from html import escape
from html.parser import HTMLParser
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, Field, model_validator

# ---------------------------------------------------------------------- mô hình
Kind = Literal["pc", "gv", "door", "board", "decor", "ban-sv", "other"]
Dir = Literal["", "dir-col", "dir-col-rev", "dir-row", "dir-row-rev"]
Decor = Literal["y-te", "cuu-hoa", "ke-giay", "tu-do", "tu-rack"]


class MapItem(BaseModel):
    id: str = Field(min_length=1, max_length=64, pattern=r"^[A-Za-z0-9_.-]+$")
    kind: Kind
    x: float = 0  # px trong nền (canvasWorkspace); CỬA: không dùng x
    y: float = 0  # px trong nền; CỬA: tính từ mép trên viền phòng (như hệ cũ — cửa nằm trong roomBounds)
    w: float | None = None
    h: float | None = None
    seat: str | None = Field(default=None, max_length=16, description="Chữ trên thẻ máy (.pc-name): số ghế / GV")
    dir: Dir = ""
    pos: int | None = Field(default=None, ge=1, le=4, description="Cửa 1–2 tường trái, 3–4 tường phải")
    decor: Decor | None = None
    label: str | None = Field(default=None, max_length=60, description="Chữ trên bảng / số bàn SV")
    group: str | None = Field(default=None, max_length=40, description="Mã cụm máy (nhóm có tên)")
    gname: str | None = Field(default=None, max_length=60, description="Tên cụm")
    code: str | None = Field(default=None, max_length=60, description="Mã thiết bị chung")
    monitor_code: str | None = Field(default=None, max_length=60)
    case_code: str | None = Field(default=None, max_length=60)
    html: str | None = Field(default=None, max_length=20000, description="Chỉ cho kind=other: giữ nguyên HTML cũ")


class Rect(BaseModel):
    x: float = 0
    y: float = 0
    w: float = 1000
    h: float = 1300


class MapModel(BaseModel):
    v: Literal[1] = 1
    canvas_w: float = Field(default=1000, gt=0, le=20000)
    bounds: Rect = Field(default_factory=Rect)
    items: list[MapItem] = Field(default_factory=list, max_length=3000)

    @model_validator(mode="after")
    def _khong_trung(self) -> MapModel:
        ids = [i.id for i in self.items]
        if len(ids) != len(set(ids)):
            raise ValueError("Trùng id phần tử sơ đồ")
        doors = [i.pos for i in self.items if i.kind == "door"]
        if len(doors) != len(set(doors)):
            raise ValueError("Trùng vị trí cửa")
        return self


PC_W, PC_H = 85, 95
GAP_IN, GAP_DESK, GAP_ROW = 20, 80, 22
START_X, START_Y = 90, 90
DESK_W = 2 * PC_W + GAP_IN


def col_to_x(col: int) -> float:
    desk, in_desk = divmod(col, 2)
    return START_X + desk * (DESK_W + GAP_DESK) + in_desk * (PC_W + GAP_IN)


def row_to_y(row: int) -> float:
    return START_Y + row * (PC_H + GAP_ROW)


# ---------------------------------------------------------------------- đọc tệp cũ
def _px(v: object) -> float | None:
    m = re.match(r"^\s*(-?\d+(?:\.\d+)?)", str(v or ""))
    return round(float(m.group(1)), 3) if m else None  # làm tròn như khi xuất → đọc-xuất-đọc không trôi số


_VOID = {"br", "img", "input", "hr", "meta", "link", "wbr", "source", "area", "col", "embed"}


class _Scan(HTMLParser):
    """Đọc MỘT phần tử gốc của trình vẽ cũ: thuộc tính gốc + các mẩu chữ/title cần thiết bên trong."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root: dict[str, str] | None = None
        self.stack: list[set[str]] = []
        self.pc_name = ""
        self.ban_sv = ""
        self.span = ""
        self.direct = ""
        self.monitor_title = ""
        self.case_title = ""

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        a = {k: (v or "") for k, v in attrs}
        cls = set(a.get("class", "").split())
        if self.root is None:
            self.root = a
        if "pc-monitor" in cls:
            self.monitor_title = a.get("title", "")
        if "pc-tower" in cls:
            self.case_title = a.get("title", "")
        if tag not in _VOID:
            self.stack.append(cls | {f"<{tag}>"})

    def handle_endtag(self, tag: str) -> None:
        if tag in _VOID:  # <br/> tự đóng: HTMLParser gọi cả endtag — không được lấy nhầm phần tử cha ra
            return
        while self.stack:  # HTML cũ có thể thiếu thẻ đóng → bỏ tới đúng thẻ đang đóng
            if f"<{tag}>" in self.stack.pop():
                break

    def handle_data(self, data: str) -> None:
        flat = set().union(*self.stack) if self.stack else set()
        if "pc-name" in flat:
            self.pc_name += data
        elif "ban-sv-num" in flat:
            self.ban_sv += data
        elif "<span>" in flat and "d-none" not in flat:
            self.span += data
        elif len(self.stack) == 1:
            self.direct += data


def _code(title: str, *prefixes: str) -> str | None:
    t = title or ""
    for p in prefixes:
        if t.startswith(p):
            t = t[len(p):]
            break
    t = t.replace("Chưa cập nhật", "").strip()
    return t or None


_DECOR_LABEL = {"y-te": "Y Tế", "cuu-hoa": "PCCC", "ke-giay": "Kệ Giày", "tu-do": "Tủ Đồ", "tu-rack": "Tủ Rack"}
_DECOR_ICON = {"y-te": "bi-plus-circle-dotted", "cuu-hoa": "bi-fire", "ke-giay": "bi-stack", "tu-do": "bi-grid-3x3-gap",
               "tu-rack": "bi-server"}
_DECOR_CLS = {"cuu-hoa": " decor-red", "tu-rack": " decor-gray"}
_DIRS = ("dir-col-rev", "dir-row-rev", "dir-col", "dir-row")


DOOR_H = 80  # .door-wrapper 40×80


def _door_default_y(pos: int, bounds_h: float) -> float:
    """Cửa chưa từng kéo (không có `top` inline) đứng theo CSS: cửa 1/3 `top: 80px`, cửa 2/4 `bottom: 185px`."""
    return 80.0 if pos in (1, 3) else round(bounds_h - 185 - DOOR_H, 3)


def _item_from_legacy(key: str, s: dict, bounds_h: float = 1300) -> MapItem:
    p = _Scan()
    p.feed(str(s.get("html") or ""))
    p.close()
    a = p.root or {}
    cls = set(a.get("class", "").split())
    iid = re.sub(r"[^A-Za-z0-9_.-]", "-", str(a.get("id") or key or "item"))[:64] or "item"
    x, y, w, h = _px(s.get("l")), _px(s.get("t")), _px(s.get("w")), _px(s.get("h"))
    base = {"id": iid, "x": x or 0, "y": y or 0, "w": w, "h": h}
    code = _code(a.get("title", ""), "Mã thiết bị: ")
    if "door-wrapper" in cls:
        m = re.search(r"door-pos-(\d)", a.get("class", ""))
        pos = int(m.group(1)) if m else 1
        return MapItem(id=f"door-{pos}", kind="door", pos=pos, y=y if y is not None else _door_default_y(pos, bounds_h))
    if "pc-node" in cls:
        is_gv = a.get("data-id") == "desk" or iid.startswith("item-desk")
        seat = (p.pc_name.strip() or a.get("data-hostname", "").strip() or ("GV" if is_gv else iid))[:16]
        d = next((c for c in _DIRS if c in cls), "")
        base.update(w=w or PC_W, h=h or PC_H)  # thiếu cỡ = cỡ CSS mặc định của .pc-node
        return MapItem(**base, kind="gv" if is_gv else "pc", seat=seat, dir=d,  # type: ignore[arg-type]
                       group=a.get("data-group") or None, gname=a.get("data-gname") or None,
                       monitor_code=_code(p.monitor_title, "Mã màn hình GV: ", "Mã màn hình: "),
                       case_code=_code(p.case_title, "Mã Case GV: ", "Mã Case: "))
    if "bang" in cls:
        return MapItem(**base, kind="board", label=(p.direct.strip() or "BẢNG KÍNH")[:60], code=code)
    if "decor-box" in cls and a.get("data-id") in _DECOR_LABEL:
        return MapItem(**base, kind="decor", decor=a["data-id"], label=(p.span.strip() or None), code=code)  # type: ignore[arg-type]
    if "ban-sv" in cls:
        return MapItem(**base, kind="ban-sv", label=p.ban_sv.strip()[:60] or None, code=code)
    return MapItem(**base, kind="other", html=str(s.get("html") or "")[:20000])


def from_legacy(obj: dict) -> MapModel:
    """Tệp sơ đồ cũ (dict) → MapModel. Bỏ lớp TRẠNG THÁI (online / err-*) — đó là màu lúc chụp, không phải thiết kế."""
    b = obj.get("bounds") or {}
    bounds = Rect(x=_px(b.get("l")) or 0, y=_px(b.get("t")) or 0, w=_px(b.get("w")) or 1000, h=_px(b.get("h")) or 1300)
    items: list[MapItem] = []
    seen: set[str] = set()
    for key, s in (obj.get("items") or {}).items():
        if not isinstance(s, dict):
            continue
        it = _item_from_legacy(str(key), s, bounds.h)
        n, base = 2, it.id
        while it.id in seen:  # hệ cũ có thể trùng id (2 máy cùng số) → đổi id, giữ cả hai
            it.id = f"{base}-{n}"[:64]
            n += 1
        seen.add(it.id)
        items.append(it)
    return MapModel(canvas_w=_px(obj.get("canvasWidth")) or bounds.w, bounds=bounds, items=items)


def read_legacy(maps_dir: str, room_name: str) -> MapModel | None:
    """Đọc tệp sơ đồ của hệ cũ — CHỈ ĐỌC. Tên tệp: '<phòng>'.replace(' ', '_') + '.json' (như hệ cũ + hệ thi)."""
    if not maps_dir:
        return None
    fp = Path(maps_dir) / (safe_name(room_name) + ".json")
    if not fp.is_file():
        return None
    try:
        return from_legacy(json.loads(fp.read_text(encoding="utf-8")))
    except (OSError, ValueError):
        return None


def safe_name(room_name: str) -> str:
    return os.path.basename(str(room_name or "").replace(" ", "_"))


# ---------------------------------------------------------------------- xuất tệp cũ (cho hệ thi)
def _f(v: float | None) -> str:
    return ("%.3f" % (v or 0)).rstrip("0").rstrip(".")


def _pxs(v: float | None) -> str:
    return f"{_f(v)}px" if v is not None else ""


def _attr(name: str, v: str | None) -> str:
    return f' {name}="{escape(v)}"' if v else ""


def _style(it: MapItem) -> str:
    return (f"inset: {_f(it.y)}px auto auto {_f(it.x)}px; width: {_f(it.w or PC_W)}px; "
            f"height: {_f(it.h or PC_H)}px; transform: none;")


def _pc_cls(it: MapItem) -> str:
    """Lớp của thẻ máy như trình vẽ cũ (KHÔNG lớp trạng thái online/err-*): GV không có `map-pc`."""
    kind = "alignable-pc" if it.kind == "gv" else "map-pc alignable-pc"
    return f"draggable-item pc-node {kind}" + (f" {it.dir}" if it.dir else "")


def _html(it: MapItem) -> str:
    if it.kind in ("pc", "gv"):
        gv = it.kind == "gv"
        cls = _pc_cls(it)
        mon = (("Mã màn hình GV: " if gv else "Mã màn hình: ") + it.monitor_code) if it.monitor_code else (
            "Mã màn hình GV: Chưa cập nhật" if gv else "")
        case = (("Mã Case GV: " if gv else "Mã Case: ") + it.case_code) if it.case_code else (
            "Mã Case GV: Chưa cập nhật" if gv else "")
        ident = ' data-id="desk"' if gv else _attr("data-hostname", it.seat)
        grp = _attr("data-group", it.group) + _attr("data-gname", it.gname)
        return (f'<div class="{cls}" id="{escape(it.id)}"{ident}{grp} style="{_style(it)}">'
                f'<div class="pc-hardware"><i class="bi bi-display pc-monitor"{_attr("title", mon)}></i>'
                f'<div class="pc-tower"{_attr("title", case)}></div></div>'
                f'<div class="pc-name">{escape(it.seat or ("GV" if gv else ""))}</div></div>')
    title = "Mã thiết bị: " + (it.code or "Chưa cập nhật")
    if it.kind == "board":
        return (f'<div class="draggable-item bang" id="{escape(it.id)}" data-id="bang"{_attr("title", title)} '
                f'style="{_style(it)}">{escape(it.label or "BẢNG KÍNH")}</div>')
    if it.kind == "decor" and it.decor:
        label = escape(it.label or _DECOR_LABEL[it.decor])
        return (f'<div class="draggable-item decor-box{_DECOR_CLS.get(it.decor, "")}" id="{escape(it.id)}" '
                f'data-id="{it.decor}"{_attr("title", title)} style="{_style(it)}">'
                f'<i class="bi {_DECOR_ICON[it.decor]}"></i><span>{label}</span></div>')
    if it.kind == "ban-sv":
        n = escape(it.label or "")
        return (f'<div class="draggable-item ban-sv" id="{escape(it.id)}" data-id="ban-sv"{_attr("title", title)} '
                f'style="{_style(it)}"><div class="ban-sv-num">{n}</div><span class="d-none">Bàn SV {n}</span></div>')
    if it.kind == "door":
        p = it.pos or 1
        side = "left" if p <= 2 else "right"
        return (f'<div class="door-wrapper door-pos-{p} swing-out-{side} draggable-item lock-x" data-id="door-{p}" '
                f'style="top: {_f(it.y)}px; bottom: auto;"><div class="door-leaf leaf-top"></div>'
                f'<div class="door-leaf leaf-bot"></div></div>')
    return it.html or ""


def to_legacy(m: MapModel) -> dict:
    """MapModel → dict đúng định dạng tệp sơ đồ cũ (đọc được bằng trình vẽ cũ VÀ sodo.js của hệ thi)."""
    items: dict[str, dict] = {}
    pc_classes: dict[str, str] = {}
    for it in m.items:
        if it.kind == "door":
            key = f"door-{it.pos or 1}"
            items[key] = {"t": _pxs(it.y), "l": "", "w": "", "h": "", "html": _html(it), "parentId": "roomBounds"}
            continue
        items[it.id] = {"t": _pxs(it.y), "l": _pxs(it.x), "w": _pxs(it.w or (PC_W if it.kind == "pc" else None)),
                        "h": _pxs(it.h or (PC_H if it.kind == "pc" else None)), "html": _html(it),
                        "parentId": "canvasWorkspace"}
        if it.kind in ("pc", "gv"):
            pc_classes[it.id] = _pc_cls(it)
    b = m.bounds
    return {"bounds": {"t": _pxs(b.y), "l": _pxs(b.x), "w": _pxs(b.w), "h": _pxs(b.h)}, "items": items,
            "pcClasses": pc_classes, "canvasWidth": _pxs(m.canvas_w), "machineGridHTML": ""}


def export_file(out_dir: str | Path, room_name: str, m: MapModel) -> Path:
    """Ghi tệp sơ đồ định dạng cũ — NGUYÊN TỬ (tệp tạm cùng thư mục + os.replace) để hệ thi không đọc phải tệp dở."""
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)
    name = safe_name(room_name)
    if not name:
        raise ValueError("Thiếu tên phòng")
    fp = d / (name + ".json")
    fd, tmp = tempfile.mkstemp(prefix=name + ".", suffix=".tmp", dir=d)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="") as f:
            f.write(json.dumps(to_legacy(m), ensure_ascii=False, separators=(",", ":")))
        os.replace(tmp, fp)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return fp


def remove_export(out_dir: str | Path, room_name: str) -> bool:
    fp = Path(out_dir) / (safe_name(room_name) + ".json")
    if fp.is_file():
        fp.unlink()
        return True
    return False


# ---------------------------------------------------------------------- Excel (openpyxl)
def _norm_seat(v: object) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():  # ô số Excel 7 → "7.0"
        s = s[:-2]
    if not s:
        return None
    return ("0" + s) if re.fullmatch(r"\d", s) else s[:16]


def read_excel(data: bytes) -> tuple[str, list[list]]:
    """Trả (loại, dữ liệu): ("grid", ma trận số ghế) hoặc ("codes", các dòng mã thiết bị) — như 2 nút nhập của hệ cũ."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — tệp hỏng / không phải xlsx
        raise ValueError("Không đọc được tệp Excel (.xlsx)") from e
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    if not rows:
        raise ValueError("Tệp Excel trống")
    head = str(rows[0][0] if rows[0] else "").strip().lower()
    if head.startswith("id"):
        return "codes", rows[1:]
    title_idx = 1 if "cột" in head else 0
    return "grid", rows[title_idx + 1:]


def apply_grid(m: MapModel, matrix: list[list]) -> tuple[MapModel, int]:
    """Thay TOÀN BỘ máy (kind=pc) bằng lưới số ghế; giữ bàn GV/cửa/bảng/đồ vật. Giãn viền phòng như hệ cũ."""
    width = max((len(r) for r in matrix), default=0)
    rows = [[_norm_seat(r[c] if c < len(r) else None) for c in range(width)] for r in matrix]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    keep = [it for it in m.items if it.kind != "pc"]
    used = {it.id for it in keep}
    pcs: list[MapItem] = []
    for r, row in enumerate(rows):
        for c, seat in enumerate(row):
            if seat is None:
                continue
            iid, n = re.sub(r"[^A-Za-z0-9_.-]", "-", f"pc-{seat}")[:60], 2
            base = iid
            while iid in used:
                iid = f"{base}-{n}"
                n += 1
            used.add(iid)
            pcs.append(MapItem(id=iid, kind="pc", x=col_to_x(c), y=row_to_y(r), w=PC_W, h=PC_H, seat=seat))
    out = m.model_copy(deep=True)
    out.items = keep + pcs
    if rows and width:
        w = max(1000, col_to_x(width - 1) + PC_W + 120)
        h = max(1300, row_to_y(len(rows) - 1) + PC_H + 300)
        out.bounds = Rect(x=out.bounds.x, y=out.bounds.y, w=w, h=h)
        out.canvas_w = w
    return out, len(pcs)


def apply_codes(m: MapModel, rows: list[list]) -> tuple[MapModel, int]:
    """Cột: ID · Tên/Vị trí · Mã thiết bị chung · Mã màn hình (PC) · Mã case (PC) — như "Nhập mã thiết bị" hệ cũ."""
    out = m.model_copy(deep=True)
    by_id = {it.id: it for it in out.items}
    n = 0
    for r in rows:
        if not r or not r[0]:
            continue
        it = by_id.get(str(r[0]).strip())
        if it is None:
            continue
        cell = [str(v).strip() if v is not None else "" for v in (list(r) + [None] * 5)[:5]]
        if it.kind in ("pc", "gv"):
            if cell[3] or cell[4]:
                it.monitor_code = cell[3] or it.monitor_code
                it.case_code = cell[4] or it.case_code
                n += 1
        elif cell[2]:
            it.code = cell[2]
            n += 1
    return out, n


def _clusters(vals: list[float], thr: float = 40) -> list[float]:
    cs: list[list[float]] = []
    for v in sorted(vals):
        if cs and abs(v - sum(cs[-1]) / len(cs[-1])) < thr:
            cs[-1].append(v)
        else:
            cs.append([v])
    return [sum(c) / len(c) for c in cs]


def _grid_index(v: float, fn, thr: float = 30, limit: int = 200) -> int | None:
    best = min(range(limit), key=lambda i: abs(fn(i) - v))
    return best if abs(fn(best) - v) <= thr else None


def _seat_grid(pcs: list[MapItem]) -> list[list[str]]:
    """Máy → lưới số ghế. Máy nằm đúng mốc lưới của hệ cũ (col_to_x / row_to_y, lệch ≤ 30px — như snap của trình vẽ)
    → suy NGƯỢC đúng chỉ số cột/hàng, GIỮ ô trống (lối đi): nhập lại ra y nguyên chỗ. Máy đặt tự do → gom cụm theo
    toạ độ (thứ tự đúng, khoảng trống không giữ được)."""
    if not pcs:
        return []
    idx = [(_grid_index(it.y, row_to_y), _grid_index(it.x, col_to_x)) for it in pcs]
    if any(r is None or c is None for r, c in idx):
        rows, cols = _clusters([it.y for it in pcs]), _clusters([it.x for it in pcs])

        def near(v: float, arr: list[float]) -> int:
            return min(range(len(arr)), key=lambda i: abs(arr[i] - v))

        idx = [(near(it.y, rows), near(it.x, cols)) for it in pcs]
    nr = max(r for r, _ in idx if r is not None) + 1
    nc = max(c for _, c in idx if c is not None) + 1
    grid = [["" for _ in range(nc)] for _ in range(nr)]
    for it, (r, c) in zip(pcs, idx, strict=True):
        if r is not None and c is not None and not grid[r][c]:
            grid[r][c] = it.seat or ""
    return grid


def to_excel(m: MapModel, room_name: str) -> bytes:
    """Sheet 1 "So_do": lưới số ghế (nhập lại được), sheet 2 "Ma_Thiet_Bi": mã thiết bị theo ID — như hệ cũ."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    pcs = [it for it in m.items if it.kind == "pc"]
    grid = _seat_grid(pcs)
    cols = grid[0] if grid else []
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "So_do"
    ws.append([f"SƠ ĐỒ PHÒNG {room_name.upper()}"])
    ws["A1"].font = Font(bold=True)
    if cols:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(cols)))
    for row in grid:
        ws.append(row)
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(horizontal="center")
    ws2 = wb.create_sheet("Ma_Thiet_Bi")
    ws2.append(["ID (Không sửa)", "Tên/Vị trí", "Mã thiết bị chung", "Mã Màn hình (Chỉ PC)", "Mã Case (Chỉ PC)"])
    for it in m.items:
        if it.kind in ("pc", "gv"):
            ws2.append([it.id, it.seat or "", "", it.monitor_code or "", it.case_code or ""])
        elif it.kind in ("board", "decor", "ban-sv"):
            name = it.label or (_DECOR_LABEL.get(it.decor, "") if it.decor else "")
            ws2.append([it.id, name, it.code or "", "", ""])
    for col, wch in zip("ABCDE", (18, 16, 22, 24, 24), strict=True):
        ws2.column_dimensions[col].width = wch
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
