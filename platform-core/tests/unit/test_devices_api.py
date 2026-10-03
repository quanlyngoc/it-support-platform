# pyright: reportOptionalSubscript=false, reportOptionalMemberAccess=false
"""API phòng + thiết bị (Sổ tay 5.1): số online tuyệt đối, máy offline không có CPU/RAM, phạm vi phòng theo policy,
Chuyển phòng / tự gỡ ép / Xoá máy (lưu trữ trước), quyền device.manage, cô lập tổ chức, realtime theo phòng."""
from __future__ import annotations

from app.legacy_adapter import contract as k
from app.rbac.service import create_policy
from app.realtime.hub import socketio
from app.tenants import service as tsvc
from app.tenants.repo import repo_for
from tests.conftest import auth
from tests.unit.conftest import tok

MAY = [("A6-301-01", "aa:bb:cc:00:30:01"), ("A6-301-02", "aa:bb:cc:00:30:02"), ("A6-405-01", "aa:bb:cc:00:40:01"),
       ("DESKTOP-XYZ", "aa:bb:cc:00:99:01")]


def _hb(w, host, mac, **kw):
    assert w["c"].post("/update", json=k.heartbeat(host, mac, **kw)).status_code == 200


def _dung_phong(w):
    """4 máy gửi nhịp tim → bộ ghi tạo phòng; rồi A6-301-02 mất nhịp tim 36 s → offline."""
    for h, m in MAY:
        _hb(w, h, m)
    w["ad"].writer_tick()
    w["clock"][0] += 36
    for h, m in MAY:
        if h != "A6-301-02":
            _hb(w, h, m)
    w["ad"].writer_tick()
    rooms = w["c"].get("/api/rooms", headers=auth(tok(w))).get_json()["items"]
    return {r["name"]: r for r in rooms}


def _dev(w, token, host):
    items = w["c"].get(f"/api/devices?q={host}", headers=auth(token)).get_json()["items"]
    return next(d for d in items if d["hostname"] == host)


def test_phong_va_thiet_bi_so_online_tuyet_doi(world):
    phong = _dung_phong(world)
    assert list(phong) == ["Phòng 301 - A6", "Phòng 405 - A6", "Phòng chờ"]  # phòng chính thức trước
    assert (phong["Phòng 301 - A6"]["total"], phong["Phòng 301 - A6"]["online"]) == (2, 1)
    assert phong["Phòng chờ"]["official"] is False and phong["Phòng 301 - A6"]["official"] is True
    t = tok(world)
    items = world["c"].get(f"/api/devices?room_id={phong['Phòng 301 - A6']['id']}", headers=auth(t)).get_json()["items"]
    on, off = (next(d for d in items if d["hostname"] == h) for h in ("A6-301-01", "A6-301-02"))
    assert on["is_online"] and on["cpu_usage"] == 12.5 and on["channel"] == "http" and on["serial"] == "ABC1234"
    assert not off["is_online"] and "cpu_usage" not in off and "ram_usage" not in off  # bài học 10.11
    r = world["c"].get(f"/api/rooms/{phong['Phòng 405 - A6']['id']}", headers=auth(t)).get_json()
    assert r["total"] == 1 and r["devices"][0]["hostname"] == "A6-405-01"
    assert world["c"].get("/api/devices?online=false", headers=auth(t)).get_json()["total"] == 1
    assert world["c"].get("/api/rooms/khong-phai-id", headers=auth(t)).status_code == 404


def test_pham_vi_phong_theo_policy(world, db):
    phong = _dung_phong(world)
    rid = phong["Phòng 301 - A6"]["id"]
    from bson import ObjectId

    pid = create_policy(db, world["tid"], "Giám thị 301", permission_codes=["user.view"],
                        room_ids=[ObjectId(rid)])["_id"]
    tsvc.create_user(db, world["tid"], "gs301", password="Gs@12345", policy_ids=[pid])
    t = tok(world, user="gs301", pw="Gs@12345")
    assert [r["name"] for r in world["c"].get("/api/rooms", headers=auth(t)).get_json()["items"]] == ["Phòng 301 - A6"]
    assert {d["hostname"] for d in world["c"].get("/api/devices", headers=auth(t)).get_json()["items"]} == {
        "A6-301-01", "A6-301-02"}
    id405 = _dev(world, tok(world), "A6-405-01")["id"]
    assert world["c"].get(f"/api/devices/{id405}", headers=auth(t)).status_code == 404
    assert world["c"].get(f"/api/rooms/{phong['Phòng 405 - A6']['id']}", headers=auth(t)).status_code == 404


def test_chuyen_phong_va_go_ep(world, db):
    _dung_phong(world)
    t, c, ad = tok(world), world["c"], world["ad"]
    d = _dev(world, t, "A6-405-01")
    r = c.post(f"/api/devices/{d['id']}/move", json={"room_name": "Phòng 301 - A6", "seat": "3"}, headers=auth(t))
    assert r.status_code == 200 and r.get_json()["moved"] == ["A6-405-01"]
    d = _dev(world, t, "A6-405-01")
    assert (d["room_name"], d["room_override"], d["move_seat"]) == ("Phòng 301 - A6", "Phòng 301 - A6", "3")
    world["clock"][0] += 61  # nhịp ghi định kỳ sau đó vẫn giữ phòng ép (hàm phân giải phòng duy nhất)
    _hb(world, "A6-405-01", "aa:bb:cc:00:40:01")
    ad.writer_tick()
    assert _dev(world, t, "A6-405-01")["room_name"] == "Phòng 301 - A6"
    # chuyển về đúng phòng theo tên máy → gỡ ép
    c.post(f"/api/devices/{d['id']}/move", json={"room_name": "Phòng 405 - A6"}, headers=auth(t))
    d = _dev(world, t, "A6-405-01")
    assert d["room_name"] == "Phòng 405 - A6" and "room_override" not in d
    assert ad.writer().room_of("A6-405-01") == "Phòng 405 - A6"
    # nhiều máy: số vị trí bị bỏ; id lạ → bỏ qua
    ids = [_dev(world, t, h)["id"] for h in ("A6-301-01", "A6-301-02")]
    r = c.post("/api/devices/move-bulk", json={"device_ids": [*ids, "0" * 24], "room_name": "Phòng chờ"},
               headers=auth(t)).get_json()
    assert sorted(r["moved"]) == ["A6-301-01", "A6-301-02"] and r["skipped"] == ["0" * 24]
    assert c.post("/api/devices/move-bulk", json={"device_ids": ["0" * 24], "room_name": "Phòng chờ"},
                  headers=auth(t)).status_code == 404
    targets = c.get("/api/rooms/move-targets", headers=auth(t)).get_json()["items"]
    assert targets[-1] == "Phòng chờ" and "Phòng 301 - A6" in targets


def test_phong_ep_tu_go_khi_ten_may_da_dung_vi_tri(world, db):
    _dung_phong(world)
    ad, w = world["ad"], world["ad"].writer()
    repo_for(db, world["tid"], "computers").update_one({"legacy_hostname": "A6-301-01"},
                                                       {"$set": {"room_override": "Phòng 301 - A6", "move_seat": "1"}})
    w.set_override("A6-301-01", "Phòng 301 - A6", "1")
    world["clock"][0] += 61
    _hb(world, "A6-301-01", "aa:bb:cc:00:30:01")
    ad.writer_tick()
    assert "A6-301-01" not in w.overrides
    doc = repo_for(db, world["tid"], "computers").find_one({"legacy_hostname": "A6-301-01"})
    assert doc is not None and "room_override" not in doc


def test_xoa_may_luu_tru_va_don_phong_ep(world, db):
    _dung_phong(world)
    t, c, ad = tok(world), world["c"], world["ad"]
    d = _dev(world, t, "A6-405-01")
    c.post(f"/api/devices/{d['id']}/move", json={"room_name": "Phòng 301 - A6"}, headers=auth(t))
    r = c.delete(f"/api/devices/{d['id']}", headers=auth(t))
    assert r.status_code == 200 and r.get_json() == {"deleted": "A6-405-01"}
    assert c.get(f"/api/devices/{d['id']}", headers=auth(t)).status_code == 404
    arch = repo_for(db, world["tid"], "computers_archive").find_one({"legacy_hostname": "A6-405-01"})
    assert arch is not None and str(arch["original_id"]) == d["id"] and arch["reason"] == "xoa tay boi admin"
    _hb(world, "A6-405-01", "aa:bb:cc:00:40:01")  # máy còn bật → nhịp tim kế tạo lại, về phòng theo TÊN (hết ép)
    ad.writer_tick()
    moi = _dev(world, t, "A6-405-01")
    assert moi["id"] != d["id"] and moi["room_name"] == "Phòng 405 - A6" and "room_override" not in moi
    assert c.delete(f"/api/devices/{d['id']}", headers=auth(t)).status_code == 404


def test_quyen_va_co_lap_to_chuc(world):
    _dung_phong(world)
    c = world["c"]
    d = _dev(world, tok(world), "A6-301-01")
    staff = tok(world, user="gv01", pw="Gv@12345")
    assert c.get("/api/rooms", headers=auth(staff)).status_code == 200  # xem: chỉ cần đăng nhập
    r = c.post(f"/api/devices/{d['id']}/move", json={"room_name": "Phòng chờ"}, headers=auth(staff))
    assert r.status_code == 403
    assert c.delete(f"/api/devices/{d['id']}", headers=auth(staff)).status_code == 403
    assert c.get("/api/rooms/move-targets", headers=auth(staff)).status_code == 403
    tb = tok(world, "BBB")
    assert c.get("/api/devices", headers=auth(tb)).get_json()["total"] == 0
    assert c.get("/api/rooms", headers=auth(tb)).get_json()["total"] == 0
    assert c.get(f"/api/devices/{d['id']}", headers=auth(tb)).status_code == 404
    assert c.delete(f"/api/devices/{d['id']}", headers=auth(tb)).status_code == 404
    assert c.get("/api/rooms").status_code == 401


def test_realtime_batch_device_update_theo_phong(world):
    phong = _dung_phong(world)
    ad, tid = world["ad"], world["tid"]
    rid = phong["Phòng 301 - A6"]["id"]
    xem_phong, xem_tong, nguoi_ngoai = k.browser(world["app"]), k.browser(world["app"]), k.browser(world["app"])
    srv = getattr(socketio, "server")  # noqa: B009
    srv.leave_room(k.sid_of(nguoi_ngoai), f"t:{tid}:devices", namespace="/desktop")  # Super Admin tự vào phòng tổng
    srv.leave_room(k.sid_of(xem_phong), f"t:{tid}:devices", namespace="/desktop")
    srv.enter_room(k.sid_of(xem_phong), f"t:{tid}:room:{rid}", namespace="/desktop")
    _hb(world, "A6-301-01", "aa:bb:cc:00:30:01")
    _hb(world, "A6-301-02", "aa:bb:cc:00:30:02")  # máy offline vừa online lại
    ad.writer_tick()
    goi = k.got(xem_phong, "batch_device_update")
    # total_online chỉ đếm PHÒNG CHÍNH THỨC (luật §8.2 hệ cũ): DESKTOP-XYZ ở Phòng chờ không tính
    assert len(goi) == 1 and goi[0]["room_stats"] == {rid: 2} and goi[0]["total_online"] == 3
    devs = {d["hostname"]: d for d in goi[0]["devices"]}
    assert devs["A6-301-02"]["state_changed"] is True and devs["A6-301-01"]["state_changed"] is False
    assert all(d["device_id"] for d in devs.values())  # máy mới tạo cũng có device_id ngay
    tong = k.got(xem_tong, "batch_device_update")
    assert tong and tong[0]["devices"] == [] and tong[0]["room_stats"][rid] == 2
    assert not k.got(nguoi_ngoai, "batch_device_update")  # không broadcast
    world["clock"][0] += 36  # cả phòng mất nhịp tim → offline: payload không kèm CPU/RAM
    ad.writer_tick()
    off = [d for p in k.got(xem_phong, "batch_device_update") for d in p["devices"]]
    assert off and all(not d["is_online"] and "cpu_usage" not in d for d in off)


def test_nhat_ky_thay_doi_may_chi_doc(world):
    c, ad, clock = world["c"], world["ad"], world["clock"]
    _hb(world, "A6-404-11", "aa:bb:cc:00:40:11", serial="DELL404")
    ad.writer_tick()
    clock[0] += 120
    ad.writer_tick()
    _hb(world, "A6-501-03", "aa:bb:cc:00:40:11", serial="DELL404")  # ổ cứng 501-03 lắp vào thân máy 404-11
    ad.writer_tick()
    t = tok(world)
    j = c.get("/api/devices/events", headers=auth(t)).get_json()
    assert [e["event"] for e in j["events"]] == ["trao_phan_cung"] and j["events"][0]["other_host"] == "A6-404-11"
    assert j["archive"][0]["legacy_hostname"] == "A6-404-11" and "original_id" not in j["archive"][0]
    assert c.get("/api/devices/events?event=doi_mac", headers=auth(t)).get_json()["events"] == []
    assert c.get("/api/devices/events?q=501-03", headers=auth(t)).get_json()["events"]
    assert c.get("/api/devices/events", headers=auth(tok(world, user="gv01", pw="Gv@12345"))).status_code == 403
    assert c.get("/api/devices/events", headers=auth(tok(world, "BBB"))).get_json()["events"] == []


def test_xuat_cau_hinh_excel_3_sheet(world):
    import io

    from openpyxl import load_workbook

    _dung_phong(world)
    r = world["c"].get("/api/reports/config-export", headers=auth(tok(world)))
    assert r.status_code == 200 and "spreadsheetml" in r.headers["Content-Type"]
    assert r.headers["Content-Disposition"].startswith("attachment; filename=Cau_Hinh_Phong_May_")
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Tổng hợp theo phòng", "Nhóm cấu hình", "Chi tiết máy"]
    chi_tiet = [row for row in wb["Chi tiết máy"].iter_rows(min_row=4, values_only=True)]
    assert {row[2]: row[3] for row in chi_tiet} == {"A6-301-01": "Online", "A6-301-02": "Offline",
                                                    "A6-405-01": "Online"}  # bỏ DESKTOP-XYZ (Phòng chờ)
    tong = [row for row in wb["Tổng hợp theo phòng"].iter_rows(min_row=4, values_only=True)]
    assert tong[-1][:3] == ("TOÀN TRƯỜNG", 3, 2)
    assert world["c"].get("/api/reports/config-export",
                          headers=auth(tok(world, user="gv01", pw="Gv@12345"))).status_code == 403
