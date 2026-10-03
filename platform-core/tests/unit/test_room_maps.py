# pyright: reportOptionalSubscript=false, reportOptionalMemberAccess=false
"""Sơ đồ phòng (Sổ tay 5.3): đọc tệp sơ đồ của hệ cũ, xuất lại ĐÚNG định dạng để hệ thi (sodo.js) đọc ra cùng
máy / cùng toạ độ / cùng hướng, nhập–xuất Excel như hệ cũ, API lưu + xuất tệp + realtime, quyền room_map.edit,
lệnh nhập sơ đồ cũ."""
from __future__ import annotations

import dataclasses
import glob
import io
import json
import os
import re
import time
from html.parser import HTMLParser

import pytest
from bson import ObjectId
from openpyxl import Workbook, load_workbook

from app.legacy_adapter import contract as k
from app.rbac.service import create_policy
from app.realtime.hub import DESKTOP_NS
from app.rooms import maps
from app.rooms.maps_cli import nhap_cu, xuat
from app.rooms.service import ensure_room
from app.server import create_app
from app.tenants import service as tsvc
from tests.conftest import TEST_SETTINGS, _make_tenant, auth, fake_ldap
from tests.unit.conftest import tok

REAL = r"D:\NetworkInventory\Maps"


# ---------------------------------------------------------------------- dữ liệu cũ tự dựng theo HTML trình vẽ cũ
def _it(t: str, left: str, w: str, h: str, html: str, parent: str = "canvasWorkspace") -> dict:
    return {"t": t, "l": left, "w": w, "h": h, "parentId": parent, "html": html}


def _pc(n: str, left: float, top: float, cls: str = "", attrs: str = "") -> dict:
    html = (f'<div class="draggable-item pc-node map-pc alignable-pc{cls}" id="pc-{n}" data-hostname="{n}"{attrs} '
            f'style="inset: {top}px auto auto {left}px; width: 85px; height: 95px; transform: none;" '
            f'data-mac="aa:bb" data-host="A6-301-{n}" title="A6-301-{n} — Online">\n'
            f'  <div class="pc-hardware">\n    <i class="bi bi-display pc-monitor" '
            f'title="Mã màn hình: MON-{n}"></i>\n    <div class="pc-tower"></div>\n  </div>\n'
            f'  <div class="pc-name">{n}</div>\n</div>')
    return _it(f"{top}px", f"{left}px", "85px", "95px", html)


def _door(p: int, top: str) -> dict:
    side = "left" if p <= 2 else "right"
    style = f' style="top: {top}; bottom: auto;"' if top else ""
    html = (f'<div class="door-wrapper door-pos-{p} swing-out-{side} draggable-item lock-x" data-id="door-{p}"'
            f'{style}><div class="door-leaf leaf-top"></div><div class="door-leaf leaf-bot"></div></div>')
    return _it(top, "", "", "", html, "roomBounds")


GV_HTML = ('<div class="draggable-item pc-node alignable-pc online" id="item-desk-1" data-id="desk" '
           'style="inset: 1072px auto auto 677.25px;" data-host="A6-301-GV"><div class="pc-hardware">'
           '<i class="bi bi-display pc-monitor" title="Mã màn hình GV: MON-GV"></i>'
           '<div class="pc-tower" title="Mã Case GV: Chưa cập nhật"></div></div>'
           '<div class="pc-name">GV</div><br/></div>')
BANG_HTML = ('<div class="draggable-item bang" id="item-bang-1" data-id="bang" title="Mã thiết bị: Chưa cập nhật" '
             'style="">BẢNG KÍNH\n<div class="map-resizer map-resizer-n"></div></div>')
YTE_HTML = ('<div class="draggable-item decor-box" id="item-y-te-1" data-id="y-te" title="Mã thiết bị: YT-01">'
            '<i class="bi bi-plus-circle-dotted"></i><span>Y Tế</span></div>')
BANSV_HTML = ('<div class="draggable-item ban-sv" id="item-ban-sv-1" data-id="ban-sv"><div class="ban-sv-num">03</div>'
              '<span class="d-none">Bàn SV 03</span></div>')


def legacy_sample() -> dict:
    items = {
        "door-1": _door(1, "1004.4px"), "door-2": _door(2, ""),
        "pc-01": _pc("01", 90, 90, " online dir-row"),
        "pc-02": _pc("02", 195, 90, " err-user machine-selected", ' data-group="g1" data-gname="Cụm A"'),
        "pc-07": _pc("07", 90, 207, " dir-row-rev"),
        "item-desk-1": _it("1072px", "677.25px", "215.562px", "108.359px", GV_HTML),
        "item-bang-1": _it("1267px", "298.6px", "400px", "35px", BANG_HTML),
        "item-y-te-1": _it("150px", "150px", "60px", "60px", YTE_HTML),
        "item-ban-sv-1": _it("300px", "500px", "105px", "115px", BANSV_HTML),
        "item-la-1": _it("10px", "10px", "20px", "20px", '<div class="draggable-item la" id="item-la-1">?</div>'),
    }
    return {"bounds": {"t": "0px", "l": "41.07264px", "w": "916.875px", "h": "1304px"}, "items": items,
            "pcClasses": {"pc-01": "draggable-item pc-node map-pc alignable-pc online dir-row"},
            "canvasWidth": "1000px", "machineGridHTML": ""}


# ---------------------------------------------------------------------- bản Python của sodo.js (hệ thi) để so
class _Sodo(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.stack: list[tuple[str, list[str], dict | None]] = []
        self.nodes: list[dict] = []

    def handle_starttag(self, tag, attrs):
        cls = (dict(attrs).get("class") or "").split()
        rec = None
        if "pc-node" in cls:
            rec = {"name": "", "dir": next((c for c in cls if re.fullmatch(r"dir-(col|row)(-rev)?", c)), "")}
            self.nodes.append(rec)
        if tag not in ("br", "img", "input"):
            self.stack.append((tag, cls, rec))

    def handle_endtag(self, tag):
        while self.stack:
            if self.stack.pop()[0] == tag:
                break

    def handle_data(self, data):
        if any("pc-name" in c for _, c, _ in self.stack):
            node = next((r for _, _, r in reversed(self.stack) if r is not None), None)
            if node is not None:
                node["name"] += data


def sodo_view(obj: dict) -> list[tuple]:
    """Đúng những gì sodo.js lấy từ tệp: (chữ .pc-name, left + i*95, top, hướng) của mọi .pc-node trong items."""
    def pf(v):
        m = re.match(r"^\s*(-?\d+(\.\d+)?)", str(v or ""))
        return float(m.group(1)) if m else 0.0
    out = []
    for sm in obj["items"].values():
        p = _Sodo()
        p.feed(sm.get("html") or "")
        for i, n in enumerate(p.nodes):
            out.append((n["name"].strip(), round(pf(sm.get("l")) + i * 95, 2), round(pf(sm.get("t")), 2), n["dir"]))
    return sorted(out)


# ---------------------------------------------------------------------- chuyển đổi
def test_doc_tep_cu_du_loai_phan_tu():
    m = maps.from_legacy(legacy_sample())
    by = {it.id: it for it in m.items}
    assert {it.kind for it in m.items} == {"door", "pc", "gv", "board", "decor", "ban-sv", "other"}
    assert (by["pc-01"].seat, by["pc-01"].dir, by["pc-01"].monitor_code) == ("01", "dir-row", "MON-01")
    assert (by["pc-02"].group, by["pc-02"].gname) == ("g1", "Cụm A")
    assert by["item-desk-1"].kind == "gv" and by["item-desk-1"].seat == "GV" and by["item-desk-1"].case_code is None
    assert by["door-1"].y == 1004.4 and by["door-2"].y == 1304 - 185 - 80  # cửa chưa kéo → vị trí CSS mặc định
    assert by["item-bang-1"].label == "BẢNG KÍNH" and by["item-bang-1"].code is None
    assert (by["item-y-te-1"].decor, by["item-y-te-1"].code) == ("y-te", "YT-01")
    assert by["item-ban-sv-1"].label == "03" and by["item-la-1"].html.startswith('<div class="draggable-item la"')
    assert m.bounds.x == 41.073 and m.canvas_w == 1000


def test_xuat_lai_he_thi_doc_ra_y_het_va_on_dinh():
    src = legacy_sample()
    m = maps.from_legacy(src)
    leg = maps.to_legacy(m)
    assert sodo_view(leg) == sodo_view(src)  # hệ thi thấy đúng máy, toạ độ, hướng như tệp gốc
    assert ("GV", 677.25, 1072.0, "") in sodo_view(leg)
    assert maps.from_legacy(json.loads(json.dumps(leg))).model_dump() == m.model_dump()
    assert set(leg) == {"bounds", "items", "pcClasses", "canvasWidth", "machineGridHTML"}
    assert leg["items"]["door-2"]["parentId"] == "roomBounds" and leg["bounds"]["l"] == "41.073px"
    assert all("online" not in c and "err-" not in c and "selected" not in c for c in leg["pcClasses"].values())
    assert leg["items"]["item-la-1"]["html"] == src["items"]["item-la-1"]["html"]  # phần tử lạ giữ nguyên


@pytest.mark.skipif(not os.path.isdir(REAL), reason="chỉ chạy trên máy có kho sơ đồ của hệ cũ")
def test_moi_so_do_that_cua_he_cu():
    files = sorted(glob.glob(os.path.join(REAL, "*.json")))
    assert files
    for f in files:
        with open(f, encoding="utf-8") as fh:  # CHỈ ĐỌC kho cũ
            src = json.load(fh)
        m = maps.from_legacy(src)
        leg = maps.to_legacy(m)
        assert sodo_view(leg) == sodo_view(src), f
        assert maps.from_legacy(json.loads(json.dumps(leg))).model_dump() == m.model_dump(), f


def test_trung_id_hoac_trung_cua_bi_tu_choi():
    with pytest.raises(ValueError):
        maps.MapModel(items=[maps.MapItem(id="a", kind="pc"), maps.MapItem(id="a", kind="pc")])
    with pytest.raises(ValueError):
        maps.MapModel(items=[maps.MapItem(id="d1", kind="door", pos=1), maps.MapItem(id="d2", kind="door", pos=1)])
    dup = legacy_sample()
    dup["items"]["pc-01-copy"] = dict(dup["items"]["pc-01"])  # hệ cũ: 2 phần tử cùng id="pc-01" → giữ cả hai
    ids = [it.id for it in maps.from_legacy(dup).items]
    assert "pc-01" in ids and "pc-01-2" in ids


# ---------------------------------------------------------------------- Excel
def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_luoi_so_ghe_va_ma_thiet_bi():
    base = maps.from_legacy(legacy_sample())
    kind, rows = maps.read_excel(_xlsx([["SƠ ĐỒ PHÒNG A6-301"], [3, 4, None, 5], ["2", 1, None, "GV2"], [None, None]]))
    assert kind == "grid"
    m, n = maps.apply_grid(base, rows)
    pcs = {it.seat: it for it in m.items if it.kind == "pc"}
    assert n == 6 and set(pcs) == {"03", "04", "05", "02", "01", "GV2"}  # 1 chữ số → thêm 0 như hệ cũ
    # Quy tắc hệ cũ: máy 85×95, 2 máy/bàn cách 20, giữa 2 bàn 80, hàng cách 22, gốc (90, 90); ô trống giữ chỗ
    assert (pcs["03"].x, pcs["03"].y) == (90, 90) and (pcs["04"].x, pcs["04"].y) == (195, 90)
    assert (pcs["05"].x, pcs["05"].y) == (90 + 270 + 105, 90) and (pcs["01"].x, pcs["01"].y) == (195, 207)
    assert (pcs["GV2"].x, pcs["GV2"].y) == (465, 207)
    assert {it.kind for it in m.items if it.kind != "pc"} == {"door", "gv", "board", "decor", "ban-sv", "other"}
    assert m.bounds.w >= 1000 and m.bounds.h >= 1300
    # Xuất → nhập lại: cùng số ghế ở cùng ô lưới
    kind2, rows2 = maps.read_excel(maps.to_excel(m, "Phòng 301 - A6"))
    m2, _ = maps.apply_grid(base, rows2)
    key = lambda mm: sorted((it.seat, it.x, it.y) for it in mm.items if it.kind == "pc")  # noqa: E731
    assert kind2 == "grid" and key(m2) == key(m)
    wb = load_workbook(io.BytesIO(maps.to_excel(m, "Phòng 301 - A6")))
    assert wb.sheetnames == ["So_do", "Ma_Thiet_Bi"] and wb["So_do"]["A1"].value == "SƠ ĐỒ PHÒNG PHÒNG 301 - A6"
    # Bảng mã thiết bị (cột A "ID…") → cập nhật mã, không đụng vị trí
    kind3, rows3 = maps.read_excel(_xlsx([["ID (Không sửa)", "Tên", "Mã chung", "Mã màn hình", "Mã case"],
                                          ["pc-01", "01", "", "MON-X", "CASE-X"],
                                          ["item-bang-1", "Bảng", "B-01", "", ""],
                                          ["khong-co", "", "Z", "", ""]]))
    m3, n3 = maps.apply_codes(base, rows3)
    by = {it.id: it for it in m3.items}
    assert kind3 == "codes" and n3 == 2
    assert (by["pc-01"].monitor_code, by["pc-01"].case_code, by["item-bang-1"].code) == ("MON-X", "CASE-X", "B-01")
    with pytest.raises(ValueError):
        maps.read_excel(b"khong phai excel")


# ---------------------------------------------------------------------- API
@pytest.fixture
def mw(db, tmp_path):
    legacy = tmp_path / "Maps_cu"
    legacy.mkdir()
    (legacy / "Phòng_301_-_A6.json").write_text(json.dumps(legacy_sample(), ensure_ascii=False), encoding="utf-8")
    t = {"P": _make_tenant(db, "PHENIKAA"), "B": _make_tenant(db, "BBB")}
    s = dataclasses.replace(TEST_SETTINGS, DATA_DIR=str(tmp_path / "data"), LEGACY_MAPS_DIR=str(legacy))
    a = create_app(db=db, settings=s, async_mode="threading", ldap_bind=fake_ldap)
    tid = t["P"]["tenant"]["_id"]
    rid = str(ensure_room(db, tid, "Phòng 301 - A6"))
    rid405 = str(ensure_room(db, tid, "Phòng 405 - A6"))
    return {"app": a, "c": a.test_client(), "tid": tid, "t": t, "rid": rid, "rid405": rid405, "tmp": tmp_path,
            "out": tmp_path / "data" / "PHENIKAA" / "Maps", "clock": [time.time()]}


def test_api_doc_cu_luu_xuat_tep_va_realtime(mw):
    c, rid = mw["c"], mw["rid"]
    t = tok(mw)
    j = c.get(f"/api/rooms/{rid}/map", headers=auth(t)).get_json()
    assert j["source"] == "legacy" and j["can_edit"] and len(j["map"]["items"]) == 10
    assert c.get(f"/api/rooms/{mw['rid405']}/map", headers=auth(t)).get_json()["source"] is None
    d = k.browser(mw["app"], t)
    assert d.emit("subscribe_room", {"room_id": rid}, namespace=DESKTOP_NS, callback=True)["ok"]
    d.get_received(DESKTOP_NS)
    m = j["map"]
    next(it for it in m["items"] if it["id"] == "pc-01")["x"] = 300
    r = c.put(f"/api/rooms/{rid}/map", json=m, headers=auth(t))
    assert r.status_code == 200
    assert r.get_json()["exported"] == "Phòng_301_-_A6.json" and not r.get_json()["export_error"]
    out = json.loads((mw["out"] / "Phòng_301_-_A6.json").read_text(encoding="utf-8"))
    assert ("01", 300.0, 90.0, "dir-row") in sodo_view(out)  # hệ thi đọc được bản mới
    assert k.got(d, "room_map_updated")[-1] == {"room_id": rid, "by": "admin"}
    j = c.get(f"/api/rooms/{rid}/map", headers=auth(t)).get_json()
    assert j["source"] == "core" and j["updated_by"] == "admin"
    # trùng id → 422, không lưu
    bad = dict(m, items=m["items"] + [m["items"][0]])
    assert c.put(f"/api/rooms/{rid}/map", json=bad, headers=auth(t)).status_code == 422
    # Xuất Excel từ bản đã lưu; tên tệp có dấu đi qua filename*
    r = c.post(f"/api/rooms/{rid}/map/export-excel", headers=auth(t))
    assert r.status_code == 200 and r.data[:2] == b"PK" and "filename*=UTF-8''So_Do_Ph%C3%B2ng" in r.headers[
        "Content-Disposition"]
    # Xoá: không đọc lại tệp hệ cũ nữa + gỡ tệp đã xuất
    assert c.delete(f"/api/rooms/{rid}/map", headers=auth(t)).get_json() == {"ok": True, "export_removed": True}
    assert c.get(f"/api/rooms/{rid}/map", headers=auth(t)).get_json()["source"] is None
    assert not (mw["out"] / "Phòng_301_-_A6.json").exists()
    assert c.post(f"/api/rooms/{rid}/map/export-excel", headers=auth(t)).status_code == 404


def test_api_nhap_excel_tra_so_do_moi_chua_luu(mw):
    c, rid, t = mw["c"], mw["rid"], tok(mw)
    data = {"file": (io.BytesIO(_xlsx([["SƠ ĐỒ"], [1, 2], [3]])), "so_do.xlsx")}
    r = c.post(f"/api/rooms/{rid}/map/import-excel", data=data, headers=auth(t), content_type="multipart/form-data")
    j = r.get_json()
    assert r.status_code == 200 and j["kind"] == "grid" and j["count"] == 3
    assert sorted(it["seat"] for it in j["map"]["items"] if it["kind"] == "pc") == ["01", "02", "03"]
    assert any(it["kind"] == "gv" for it in j["map"]["items"])  # nền là tệp cũ → giữ bàn GV
    assert c.get(f"/api/rooms/{rid}/map", headers=auth(t)).get_json()["source"] == "legacy"  # CHƯA lưu
    sent = {"file": (io.BytesIO(_xlsx([["SƠ ĐỒ"], [9]])), "x.xlsx"), "map": json.dumps(maps.MapModel().model_dump())}
    j = c.post(f"/api/rooms/{rid}/map/import-excel", data=sent, headers=auth(t),
               content_type="multipart/form-data").get_json()
    assert [it["seat"] for it in j["map"]["items"]] == ["09"]  # nền là sơ đồ ĐANG SỬA gửi kèm
    assert c.post(f"/api/rooms/{rid}/map/import-excel", headers=auth(t)).status_code == 422
    bad = {"file": (io.BytesIO(b"rac"), "x.xlsx")}
    assert c.post(f"/api/rooms/{rid}/map/import-excel", data=bad, headers=auth(t),
                  content_type="multipart/form-data").status_code == 422


def test_api_quyen_pham_vi_va_xuat_loi(mw, db):
    c, rid, t = mw["c"], mw["rid"], tok(mw)
    staff = tok(mw, user="gv01", pw="Gv@12345")  # policy chỉ xem
    j = c.get(f"/api/rooms/{rid}/map", headers=auth(staff)).get_json()
    assert j["can_edit"] is False
    assert c.put(f"/api/rooms/{rid}/map", json=j["map"], headers=auth(staff)).status_code == 403
    assert c.delete(f"/api/rooms/{rid}/map", headers=auth(staff)).status_code == 403
    pid = create_policy(db, mw["tid"], "GS 405", permission_codes=["room_map.edit"],
                        room_ids=[ObjectId(mw["rid405"])])["_id"]
    tsvc.create_user(db, mw["tid"], "gs405", password="Gs@12345", policy_ids=[pid])
    g = tok(mw, user="gs405", pw="Gs@12345")
    assert c.get(f"/api/rooms/{rid}/map", headers=auth(g)).status_code == 404  # ngoài phạm vi phòng
    assert c.put(f"/api/rooms/{rid}/map", json=j["map"], headers=auth(g)).status_code == 404
    tb = tok(mw, tenant="BBB")
    assert c.get(f"/api/rooms/{rid}/map", headers=auth(tb)).status_code == 404  # tổ chức khác
    assert c.get("/api/rooms/khong-phai-id/map", headers=auth(t)).status_code == 404
    # Thư mục xuất hỏng → vẫn LƯU vào Core, báo lỗi xuất để người sửa biết hệ thi chưa nhận bản mới
    blocker = mw["tmp"] / "la-tep"
    blocker.write_text("x", encoding="utf-8")
    mw["app"].config["ITM_SETTINGS"] = dataclasses.replace(mw["app"].config["ITM_SETTINGS"],
                                                           MAPS_EXPORT_DIR=str(blocker / "con"))
    r = c.put(f"/api/rooms/{rid}/map", json=j["map"], headers=auth(t)).get_json()
    assert r["ok"] and r["exported"] is None and r["export_error"]
    assert c.get(f"/api/rooms/{rid}/map", headers=auth(t)).get_json()["source"] == "core"


def test_lenh_nhap_cu_va_xuat(mw, db):
    tenant = mw["t"]["P"]["tenant"]
    src = mw["tmp"] / "Maps_cu"
    assert nhap_cu(db, tenant, src, ghi_de=False, thu=True) == 1  # chạy thử: không ghi
    c, rid, t = mw["c"], mw["rid"], tok(mw)
    assert c.get(f"/api/rooms/{rid}/map", headers=auth(t)).get_json()["source"] == "legacy"
    assert nhap_cu(db, tenant, src, ghi_de=False, thu=False) == 1
    assert c.get(f"/api/rooms/{rid}/map", headers=auth(t)).get_json()["source"] == "core"
    assert nhap_cu(db, tenant, src, ghi_de=False, thu=False) == 0  # đã có trong Core → giữ nguyên
    out = mw["tmp"] / "xuat"
    assert xuat(db, tenant, out) == 1
    got = json.loads((out / "Phòng_301_-_A6.json").read_text(encoding="utf-8"))
    assert sodo_view(got) == sodo_view(legacy_sample())
