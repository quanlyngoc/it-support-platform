# pyright: reportOptionalSubscript=false, reportOptionalMemberAccess=false
"""Quy trình sự cố (Sổ tay 8.2): báo hỏng → bổ sung → ghi chú IT → xử lý từng dòng → xong; một phiếu USER mở/máy;
sự cố hệ thống không bổ sung được; quyền incident.handle / report.export; phạm vi phòng; kênh cảnh báo email + Telegram
(bí mật mã hoá, chỉ trả mặt nạ); log máy trạm; xuất Excel."""
from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

import pytest
from bson import ObjectId
from openpyxl import load_workbook

from app.integrations import alerts
from app.rbac.service import create_policy
from app.tenants import service as tsvc
from app.tenants.repo import SettingRepo, repo_for
from tests.conftest import auth
from tests.unit.conftest import tok


@pytest.fixture
def sent(monkeypatch):
    """Kênh cảnh báo giả: gửi ĐỒNG BỘ, ghi lại thay vì gọi Telegram / SMTP thật."""
    out: dict[str, list] = {"telegram": [], "email": []}
    monkeypatch.setattr(alerts, "_spawn", lambda fn, *a: fn(*a))
    monkeypatch.setattr(alerts, "_post_telegram", lambda token, chat, text: out["telegram"].append((token, chat, text)))
    monkeypatch.setattr(alerts, "_smtp_send", lambda host, port, sender, pw, rcpts, subject, body: out["email"].append(
        (host, port, sender, pw, rcpts, subject, body)))
    return out


def _may(db, w, host: str, room: str) -> str:
    tid = w["tid"]
    rr = repo_for(db, tid, "rooms")
    r = rr.find_one({"name": room}) or {"_id": rr.insert_one({"name": room}).inserted_id}
    return str(repo_for(db, tid, "computers").insert_one(
        {"hostname": host, "legacy_hostname": host, "room_name": room, "room_id": r["_id"]}).inserted_id)


def _cau_hinh(c, t):
    r = c.put("/api/settings/alerts", headers=auth(t), json={
        "telegram_bot_token": "123:ABC", "telegram_chat_id": "-100", "sender_email": "it@x.vn",
        "receiver_email": "a@x.vn; b@x.vn\na@x.vn", "gmail_app_password": "app-pass-16"})
    assert r.status_code == 200, r.get_json()
    return r.get_json()


def test_quy_trinh_bao_hong_den_xong(world, db, sent):
    c, t = world["c"], tok(world)
    _cau_hinh(c, t)
    d301 = _may(db, world, "A6-301-05", "Phòng 301 - A6")

    r = c.post("/api/incidents", headers=auth(t), json={"device_id": d301, "description": "  Máy không\nlên mạng "})
    assert r.status_code == 201, r.get_json()
    inc = r.get_json()["incident"]
    iid = inc["id"]
    assert inc["status"] == "PENDING" and inc["source"] == "USER" and inc["reporter_name"] == "admin"
    assert inc["room_name"] == "Phòng 301 - A6" and inc["hostname"] == "A6-301-05"
    assert inc["title"] == "Máy không lên mạng"
    assert [(i["n"], i["text"], i["done"]) for i in inc["items"]] == [(1, "Máy không lên mạng", False)]
    assert inc["description"].startswith("1. ") and inc["description"].endswith(". Máy không lên mạng")
    # cảnh báo: Telegram trước, email tới MỌI người nhận (bỏ trùng)
    assert len(sent["telegram"]) == 1 and sent["telegram"][0][:2] == ("123:ABC", "-100")
    assert "A6-301-05" in sent["telegram"][0][2] and "Máy không lên mạng" in sent["telegram"][0][2]
    assert sent["email"][0][4] == ["a@x.vn", "b@x.vn"] and "A6-301-05" in sent["email"][0][5]

    # báo tiếp khi phiếu còn mở → BỔ SUNG vào phiếu đó, không tạo phiếu mới, không gửi cảnh báo lần nữa
    r = c.post("/api/incidents", headers=auth(t), json={"device_id": d301, "description": "Chuột hỏng"})
    assert r.status_code == 200 and r.get_json()["appended"] is True and r.get_json()["incident"]["id"] == iid
    r = c.post(f"/api/incidents/{iid}/append", headers=auth(t), json={"text": "Màn hình nhấp nháy"})
    assert [i["n"] for i in r.get_json()["incident"]["items"]] == [1, 2, 3]
    assert len(sent["telegram"]) == 1
    assert repo_for(db, world["tid"], "incidents").count({}) == 1

    # ghi chú IT → Đang xử lý, người ghi là người xử lý
    r = c.post(f"/api/incidents/{iid}/note", headers=auth(t), json={"text": "Đã nhận, chờ linh kiện"})
    j = r.get_json()["incident"]
    assert j["status"] == "PROCESSING" and j["resolver_name"] == "admin" and j["it_notes"][0]["lines"] == []

    # xử lý một phần (dòng 1, 3) → vẫn Đang xử lý
    r = c.post(f"/api/incidents/{iid}/resolve", headers=auth(t), json={"text": "Cắm lại dây mạng", "lines": [1, 3, 9]})
    j = r.get_json()["incident"]
    assert j["status"] == "PROCESSING" and [i["done"] for i in j["items"]] == [True, False, True]
    assert j["it_notes"][-1]["lines"] == [1, 3] and j["items"][0]["done_by"] == "admin"
    assert c.post(f"/api/incidents/{iid}/resolve", headers=auth(t),
                  json={"text": "x", "lines": [1]}).status_code == 400  # dòng đã xong rồi
    # xử lý phần còn lại (không chọn dòng = cả phiếu) → Xong
    j = c.post(f"/api/incidents/{iid}/resolve", headers=auth(t), json={"text": "Thay chuột"}).get_json()["incident"]
    assert j["status"] == "DONE" and all(i["done"] for i in j["items"]) and j["resolved_at"]
    assert j["it_notes"][-1]["lines"] == [2] and j["rev"] == 6
    assert c.post(f"/api/incidents/{iid}/append", headers=auth(t), json={"text": "y"}).status_code == 409
    assert c.post(f"/api/incidents/{iid}/resolve", headers=auth(t), json={"text": "z"}).status_code == 409
    # phiếu đã xong → báo hỏng lần sau là phiếu MỚI
    r = c.post("/api/incidents", headers=auth(t), json={"device_id": d301, "description": "Lại hỏng"})
    assert r.status_code == 201 and r.get_json()["incident"]["id"] != iid and len(sent["telegram"]) == 2
    # chi tiết + tìm theo mô tả / người xử lý
    assert c.get(f"/api/incidents/{iid}", headers=auth(t)).get_json()["status"] == "DONE"
    assert c.get("/api/incidents?q=chuột", headers=auth(t)).get_json()["total"] == 1
    assert c.get("/api/incidents?source=USER&status=DONE", headers=auth(t)).get_json()["total"] == 1
    # nội dung rỗng / id lạ
    assert c.post("/api/incidents", headers=auth(t), json={"device_id": d301, "description": "   "}).status_code == 400
    assert c.post(f"/api/incidents/{ObjectId()}/append", headers=auth(t), json={"text": "a"}).status_code == 404
    assert c.get("/api/incidents/khong-phai-id", headers=auth(t)).status_code == 404


def test_su_co_he_thong_va_canh_bao_phan_cung(world, db, sent):
    c, t = world["c"], tok(world)
    _cau_hinh(c, t)
    iid = str(repo_for(db, world["tid"], "incidents").insert_one({  # dạng bộ ghi nền tạo (không có rev, không items)
        "hostname": "A6-405-01", "room_name": "Phòng 405 - A6", "title": "MẤT RAM", "category": "hardware_ram",
        "description": "MẤT RAM: 16GB -> 8GB", "status": "PENDING", "source": "SYSTEM_WATCHDOG",
        "reporter_name": "SYSTEM_WATCHDOG", "created_at": datetime.now(UTC)}).inserted_id)
    assert c.post(f"/api/incidents/{iid}/append", headers=auth(t), json={"text": "a"}).status_code == 409
    j = c.post(f"/api/incidents/{iid}/resolve", headers=auth(t), json={"text": "Lắp lại RAM", "lines": [1]}).get_json()
    assert j["incident"]["status"] == "DONE" and j["incident"]["resolver_name"] == "admin"
    # bộ ghi nền phát hiện sự cố phần cứng → cảnh báo nguồn "HỆ THỐNG TỰ ĐỘNG"
    world["ad"]._on_incident({"hostname": "A6-405-02", "room_name": "Phòng 405 - A6", "title": "HẾT PIN CMOS",
                              "description": "CẢNH BÁO: HỆ THỐNG HẾT PIN CMOS"})
    assert "HỆ THỐNG TỰ ĐỘNG" in sent["telegram"][-1][2] and "HẾT PIN CMOS" in sent["email"][-1][6]


def test_quyen_va_pham_vi_phong(world, db, sent):
    c, ta = world["c"], tok(world)
    d301 = _may(db, world, "A6-301-05", "Phòng 301 - A6")
    d405 = _may(db, world, "A6-405-05", "Phòng 405 - A6")
    r301 = repo_for(db, world["tid"], "rooms").find_one({"name": "Phòng 301 - A6"})["_id"]
    for name, perms in (("gs301", ["device.control"]), ("it301", ["incident.handle", "report.export"])):
        pid = create_policy(db, world["tid"], f"P {name}", permission_codes=perms, room_ids=[r301])["_id"]
        tsvc.create_user(db, world["tid"], name, password="Gs@12345", policy_ids=[pid])
    tg = tok(world, user="gs301", pw="Gs@12345")
    ti = tok(world, user="it301", pw="Gs@12345")
    # phía phòng báo hỏng được máy phòng mình, không được máy phòng khác
    r = c.post("/api/incidents", headers=auth(tg), json={"device_id": d301, "description": "Hỏng bàn phím"})
    assert r.status_code == 201
    i301 = r.get_json()["incident"]["id"]
    assert c.post("/api/incidents", headers=auth(tg), json={"device_id": d405, "description": "x"}).status_code == 404
    i405 = c.post("/api/incidents", headers=auth(ta), json={"device_id": d405, "description": "x"}).get_json()[
        "incident"]["id"]
    assert c.get(f"/api/incidents/{i405}", headers=auth(tg)).status_code == 404
    assert c.post(f"/api/incidents/{i405}/append", headers=auth(tg), json={"text": "y"}).status_code == 404
    # không có incident.handle → không ghi chú / xử lý / xem log / xuất Excel
    for path in (f"/api/incidents/{i301}/note", f"/api/incidents/{i301}/resolve"):
        assert c.post(path, headers=auth(tg), json={"text": "a"}).status_code == 403
    assert c.get("/api/client-logs", headers=auth(tg)).status_code == 403
    assert c.get("/api/incidents/export", headers=auth(tg)).status_code == 403
    assert c.get("/api/settings/alerts", headers=auth(tg)).status_code == 403
    # IT của phòng 301: xử lý được phiếu 301, phiếu 405 ngoài phạm vi
    assert c.post(f"/api/incidents/{i301}/resolve", headers=auth(ti), json={"text": "Thay"}).status_code == 200
    assert c.post(f"/api/incidents/{i405}/resolve", headers=auth(ti), json={"text": "Thay"}).status_code == 404
    # tổ chức khác không thấy, không đụng được
    tb = tok(world, tenant="BBB")
    assert c.get(f"/api/incidents/{i405}", headers=auth(tb)).status_code == 404
    assert c.post("/api/incidents", headers=auth(tb), json={"device_id": d301, "description": "x"}).status_code == 404
    assert c.post("/api/incidents", json={"device_id": d301, "description": "x"}).status_code == 401


def test_cau_hinh_kenh_canh_bao(world, db, sent):
    c, t = world["c"], tok(world)
    j = c.get("/api/settings/alerts", headers=auth(t)).get_json()
    assert j["smtp_server"] == "smtp.gmail.com" and j["smtp_port"] == "587" and not j["telegram_ready"]
    j = _cau_hinh(c, t)
    assert j["telegram_ready"] and j["email_ready"] and j["gmail_app_password_mask"] == "*" * 11
    assert "gmail_app_password" not in j and "telegram_bot_token" not in j
    raw = SettingRepo(db, world["tid"]).find_one({"key": alerts.KEY})["value"]
    assert "app-pass-16" not in str(raw) and raw["gmail_app_password_enc"].startswith("v1:")  # lưu MÃ HOÁ
    # để trống bí mật = giữ nguyên; clear_ = xoá
    j = c.put("/api/settings/alerts", headers=auth(t), json={"telegram_chat_id": "-200", "telegram_bot_token": ""}
              ).get_json()
    assert j["has_telegram_bot_token"] and j["telegram_chat_id"] == "-200"
    r = c.post("/api/settings/alerts/test", headers=auth(t), json={"channel": "telegram"})
    assert r.get_json()["ok"] and sent["telegram"][-1][:2] == ("123:ABC", "-200")
    r = c.post("/api/settings/alerts/test", headers=auth(t), json={"channel": "email", "receiver_email": "z@x.vn"})
    assert r.get_json()["ok"] and sent["email"][-1][4] == ["z@x.vn"]  # dùng giá trị đang gõ, chưa lưu
    j = c.put("/api/settings/alerts", headers=auth(t), json={"clear_telegram_bot_token": True}).get_json()
    assert not j["has_telegram_bot_token"] and not j["telegram_ready"]
    r = c.post("/api/settings/alerts/test", headers=auth(t), json={"channel": "telegram"}).get_json()
    assert r["ok"] is False and "Telegram" in r["message"]
    assert c.put("/api/settings/alerts", headers=auth(t), json={"smtp_port": "abc"}).status_code == 400

    def loi(*_a):
        raise OSError("mạng hỏng")

    alerts._smtp_send = loi  # type: ignore[assignment] — monkeypatch fixture sẽ khôi phục
    r = c.post("/api/settings/alerts/test", headers=auth(t), json={"channel": "email"}).get_json()
    assert r["ok"] is False and "mạng hỏng" in r["message"]
    # tổ chức khác: cấu hình riêng
    assert not c.get("/api/settings/alerts", headers=auth(tok(world, tenant="BBB"))).get_json()["email_ready"]


def test_log_may_tram_va_xuat_excel(world, db, sent):
    c, t = world["c"], tok(world)
    for host, err in (("A6-301-05", "WinError 5"), ("A6-405-01", "timeout")):
        r = c.post("/client_log", json={"mac": "AA-BB-CC-00-00-01", "hostname": host, "context": "update",
                                        "error": err}, environ_base={"REMOTE_ADDR": "10.21.30.5"})
        assert r.status_code == 200
    j = c.get("/api/client-logs", headers=auth(t)).get_json()
    assert j["total"] == 2 and j["items"][0]["hostname"] == "A6-405-01"  # mới nhất trước
    assert c.get("/api/client-logs?hostname=a6-301-05", headers=auth(t)).get_json()["items"][0]["error"] == "WinError 5"
    assert c.get("/api/client-logs?q=timeout", headers=auth(t)).get_json()["total"] == 1

    d301 = _may(db, world, "A6-301-05", "Phòng 301 - A6")
    iid = c.post("/api/incidents", headers=auth(t), json={"device_id": d301, "description": "Hỏng nguồn"}).get_json()[
        "incident"]["id"]
    c.post(f"/api/incidents/{iid}/resolve", headers=auth(t), json={"text": "Thay nguồn"})
    r = c.get("/api/incidents/export", headers=auth(t))
    assert r.status_code == 200 and r.headers["Content-Type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(BytesIO(r.data))["Sự cố"]
    assert [x.value for x in ws[3]][:5] == ["STT", "Thời gian", "Phòng", "Máy", "Người báo"]
    row = [x.value for x in ws[4]]
    assert row[3] == "A6-301-05" and row[6] == "admin" and row[8] == "Đã xong" and "Thay nguồn" in str(row[7])
    assert c.get("/api/incidents/export?status=PENDING", headers=auth(t)).status_code == 200
