# pyright: reportOptionalSubscript=false
"""Báo cáo (Sổ tay 8.2): thống kê sử dụng năm/tháng theo phòng/máy (gồm máy chưa dùng, máy đã xoá), chuỗi theo ngày /
tháng, báo cáo một ngày (RAM + lần đổi RAM), lịch sử RAM / ổ cứng, xuất Excel; phạm vi phòng; cô lập tổ chức."""
from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

from bson import ObjectId
from openpyxl import load_workbook

from app.rbac.service import create_policy
from app.tenants import service as tsvc
from app.tenants.repo import repo_for
from tests.conftest import auth
from tests.unit.conftest import tok


def _du_lieu(db, w):
    tid = w["tid"]
    rooms = {n: repo_for(db, tid, "rooms").insert_one({"name": n}).inserted_id
             for n in ("Phòng 301 - A6", "Phòng 405 - A6")}
    comp = repo_for(db, tid, "computers")
    ids = {}
    for host, room, ram in (("A6-301-01", "Phòng 301 - A6", "8GB"), ("A6-301-02", "Phòng 301 - A6", "16GB"),
                            ("A6-405-01", "Phòng 405 - A6", "8GB"), ("A6-301-10", "Phòng 301 - A6", "8GB")):
        ids[host] = comp.insert_one({"hostname": host, "legacy_hostname": host, "room_name": room,
                                     "room_id": rooms[room], "ram": ram, "mac": f"AA-{host}"}).inserted_id
    us = repo_for(db, tid, "usage_stats")
    for host, day, mins in (("A6-301-01", "2026-09-30", 60), ("A6-301-01", "2026-10-01", 120),
                            ("A6-301-02", "2026-10-01", 30), ("A6-405-01", "2026-10-02", 600),
                            ("A6-301-01", "2025-12-31", 999)):
        us.insert_one({"device_id": ids[host], "date": day, "online_minutes": mins, "hostname": host,
                       "room_name": "Phòng 405 - A6" if host.startswith("A6-405") else "Phòng 301 - A6"})
    us.insert_one({"device_id": ObjectId(), "date": "2026-10-03", "online_minutes": 45, "hostname": "A6-301-99",
                   "room_name": "Phòng 301 - A6"})  # máy đã bị xoá tay
    inc = repo_for(db, tid, "incidents")
    for i, (host, cat, desc) in enumerate((("A6-301-02", "hardware_ram", "THÊM RAM: 8GB -> 16GB"),
                                           ("A6-405-01", "hardware_disk", "MẤT Ổ CỨNG: 2 ổ (256GB) -> 1 ổ (128GB)"),
                                           ("A6-405-01", "cmos", "CẢNH BÁO: HỆ THỐNG HẾT PIN CMOS"))):
        inc.insert_one({"hostname": host, "room_name": "Phòng 405 - A6" if "405" in host else "Phòng 301 - A6",
                        "category": cat, "title": desc.split(":")[0], "description": desc, "status": "PENDING",
                        "source": "SYSTEM_WATCHDOG", "created_at": datetime(2026, 10, 1, 8, i, tzinfo=UTC)})
    return rooms, ids


def test_thong_ke_su_dung(world, db):
    rooms, _ = _du_lieu(db, world)
    c, t = world["c"], tok(world)
    j = c.get("/api/reports/usage?year=2026", headers=auth(t)).get_json()
    assert j["total_minutes"] == 60 + 120 + 30 + 600 + 45 and j["machines"] == 5 and j["used"] == 4
    r301 = j["rooms"][0]
    assert r301["room_name"] == "Phòng 301 - A6" and r301["minutes"] == 255
    assert [(m["hostname"], m["minutes"], m["days"]) for m in r301["machines"]] == [
        ("A6-301-01", 180, 2), ("A6-301-02", 30, 1), ("A6-301-10", 0, 0), ("A6-301-99", 45, 1)]  # sắp tự nhiên
    assert r301["machines"][-1]["deleted"] is True
    j = c.get("/api/reports/usage?year=2026&month=9", headers=auth(t)).get_json()
    assert j["total_minutes"] == 60
    j = c.get(f"/api/reports/usage?year=2026&room_id={rooms['Phòng 405 - A6']}", headers=auth(t)).get_json()
    assert [r["room_name"] for r in j["rooms"]] == ["Phòng 405 - A6"] and j["total_minutes"] == 600
    assert c.get("/api/reports/usage?year=1999", headers=auth(t)).status_code == 422
    # chuỗi theo ngày / tháng
    j = c.get("/api/reports/usage/series?group=day&date_from=2026-09-30&date_to=2026-10-02",
              headers=auth(t)).get_json()
    assert [(p["period"], p["minutes"], p["machines"]) for p in j["items"]] == [
        ("2026-09-30", 60, 1), ("2026-10-01", 150, 2), ("2026-10-02", 600, 1)]
    j = c.get("/api/reports/usage/series?group=month&date_from=2025-12-01&date_to=2026-10-31",
              headers=auth(t)).get_json()
    assert [(p["period"], p["minutes"], p["machines"]) for p in j["items"]] == [
        ("2025-12", 999, 1), ("2026-09", 60, 1), ("2026-10", 795, 4)]
    assert c.get("/api/reports/usage/series?date_from=2026-10-02&date_to=2026-10-01",
                 headers=auth(t)).status_code == 400
    assert "items" in c.get("/api/reports/usage/series", headers=auth(t)).get_json()  # mặc định 30 ngày gần nhất
    # báo cáo một ngày: RAM + lần cuối RAM đổi
    j = c.get("/api/reports/usage/day?day=2026-10-01", headers=auth(t)).get_json()
    assert j["total"] == 2 and j["minutes"] == 150
    m = {i["hostname"]: i for i in j["items"]}
    assert m["A6-301-02"]["ram"] == "16GB" and m["A6-301-02"]["ram_changed_at"].startswith("2026-10-01T08:00")
    assert m["A6-301-01"]["ram_changed_at"] is None
    # lịch sử RAM / ổ cứng (không gồm CMOS) kèm MAC, tách cũ → mới
    j = c.get("/api/reports/hardware-history", headers=auth(t)).get_json()
    assert j["total"] == 2
    h = j["items"][0]
    assert (h["hostname"], h["part"], h["event_type"], h["old_value"], h["new_value"], h["mac"]) == (
        "A6-405-01", "Ổ cứng", "MẤT Ổ CỨNG", "2 ổ (256GB)", "1 ổ (128GB)", "AA-A6-405-01")
    assert c.get("/api/reports/hardware-history?q=a6-301", headers=auth(t)).get_json()["items"][0]["part"] == "RAM"
    # tổ chức khác: không thấy gì
    tb = tok(world, tenant="BBB")
    assert c.get("/api/reports/usage?year=2026", headers=auth(tb)).get_json()["total_minutes"] == 0
    assert c.get("/api/reports/hardware-history", headers=auth(tb)).get_json()["total"] == 0
    assert c.get("/api/reports/usage").status_code == 401


def test_pham_vi_phong_va_xuat_excel(world, db):
    rooms, _ = _du_lieu(db, world)
    c, t = world["c"], tok(world)
    pid = create_policy(db, world["tid"], "GS 301", permission_codes=["device.control"],
                        room_ids=[rooms["Phòng 301 - A6"]])["_id"]
    tsvc.create_user(db, world["tid"], "gs301", password="Gs@12345", policy_ids=[pid])
    tg = tok(world, user="gs301", pw="Gs@12345")
    j = c.get("/api/reports/usage?year=2026", headers=auth(tg)).get_json()
    assert [r["room_name"] for r in j["rooms"]] == ["Phòng 301 - A6"] and j["total_minutes"] == 210  # không máy đã xoá
    assert c.get(f"/api/reports/usage?year=2026&room_id={rooms['Phòng 405 - A6']}",
                 headers=auth(tg)).get_json()["total_minutes"] == 0
    j = c.get("/api/reports/usage/series?group=day&date_from=2026-10-01&date_to=2026-10-03",
              headers=auth(tg)).get_json()
    assert [p["period"] for p in j["items"]] == ["2026-10-01"]
    assert c.get("/api/reports/hardware-history", headers=auth(tg)).get_json()["total"] == 1
    assert c.get("/api/reports/usage/export?year=2026", headers=auth(tg)).status_code == 403

    r = c.get("/api/reports/usage/export?year=2026", headers=auth(t))
    assert r.status_code == 200 and "Thong_Ke_Su_Dung_2026.xlsx" in r.headers["Content-Disposition"]
    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames == ["Theo máy", "Theo phòng", "Theo tháng"]
    rows = [[x.value for x in row] for row in wb["Theo máy"].iter_rows(min_row=4)]
    assert rows[0] == [1, "Phòng 301 - A6", "A6-301-01", 180, 3.0, 2]
    assert rows[3][2] == "A6-301-99 (đã xoá)" and rows[-1][:5] == [None, "TỔNG", None, 855, 14.2]
    thang = [[x.value for x in row] for row in wb["Theo tháng"].iter_rows(min_row=4)]
    assert thang == [["2026-09", 1.0, 1], ["2026-10", 13.2, 4]]
    r = c.get("/api/reports/usage/export?year=2026&month=10", headers=auth(t))
    assert load_workbook(BytesIO(r.data)).sheetnames == ["Theo máy", "Theo phòng"]
